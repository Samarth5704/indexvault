"""API: series, load jobs + SSE, CSV import, cache, ticker check, export."""
import io
import json
import zipfile

import numpy as np
import pandas as pd
import pytest

from conftest import assert_api_conventions
from indexvault import data

NIFTY = "%5ENSEI"  # ^NSEI, URL-encoded for paths


def as_frame(payload):
    cols = payload["columns"]
    df = pd.DataFrame(payload["rows"], columns=cols).set_index(cols[0])
    df.index = pd.to_datetime(df.index)
    return df


# --------------------------------------------------------------------------- #
# Series
# --------------------------------------------------------------------------- #
def test_series_default_columns(client):
    r = client.get(f"/api/series/{NIFTY}?period=1Y")
    assert r.status_code == 200
    body = r.json()
    assert_api_conventions(body)
    assert body["name"] == "NIFTY 50" and body["freq"] == "Daily" and body["source"] == "demo"
    assert body["columns"] == ["date", "open", "high", "low", "close", "adj_close", "volume", "return"]
    assert body["meta"]["has_volume"] is False
    assert any("Synthetic" in c for c in body["meta"]["caveats"])
    df = as_frame(body)
    # first row = the trailing-return base: last close on/before (last date - 1Y)
    assert df.index[0] <= df.index[-1] - pd.DateOffset(years=1) < df.index[1]
    assert df["return"].iloc[5] == pytest.approx(df["close"].iloc[5] / df["close"].iloc[4] - 1)  # decimal


def test_series_derived_columns_are_warmed_up(client):
    r = client.get(f"/api/series/{NIFTY}?start=2020-01-01&end=2020-12-31&freq=Weekly"
                   "&columns=close,sma_20,ema_10,rsi_14,vol_26,log_return,drawdown,rebased")
    df = as_frame(r.json())
    assert df.index[0] >= pd.Timestamp("2020-01-01")
    assert df[["sma_20", "ema_10", "rsi_14", "vol_26", "log_return"]].iloc[0].notna().all()
    assert df["rebased"].iloc[0] == pytest.approx(100)
    assert df["drawdown"].max() == 0 and df["drawdown"].min() < 0
    assert df["rsi_14"].between(0, 100).all()
    full = data.resample(data.get_data("^NSEI", "2015-01-01", "2020-12-31", source="demo"), "Weekly")
    assert df["sma_20"].iloc[-1] == pytest.approx(full["Close"].rolling(20).mean().loc[df.index[-1]])


@pytest.mark.parametrize("query,code", [
    ("columns=close,sma_1", 400), ("columns=bogus", 400), ("period=7Q", 400),
    ("start=2021-01-01&end=2020-01-01", 400), ("freq=Hourly", 422), ("source=nope", 400),
])
def test_series_bad_requests(client, query, code):
    r = client.get(f"/api/series/{NIFTY}?{query}")
    assert r.status_code == code and "error" in r.json()


def test_series_csv_source_missing(client):
    r = client.get("/api/series/NOPE?source=csv")
    assert r.status_code == 404 and r.json()["error"]["code"] == "not_found"


def test_source_error_is_502(client, monkeypatch):
    def broken(*_):
        raise ConnectionError("network down")
    monkeypatch.setitem(data.SOURCES, "demo", broken)
    r = client.get("/api/series/BROKEN?period=1Y")
    assert r.status_code == 502
    assert r.json()["error"]["detail"]["reason"] == "network down"


def test_series_memo_skips_reloading(client, monkeypatch):
    client.get(f"/api/series/{NIFTY}?period=1Y")
    calls = []
    monkeypatch.setattr(data, "load_cached", lambda *a, **k: calls.append(a) or (None, {}))
    assert client.get(f"/api/series/{NIFTY}?period=1Y").status_code == 200
    assert calls == []  # served from the in-memory frame cache


# --------------------------------------------------------------------------- #
# Jobs
# --------------------------------------------------------------------------- #
def test_load_job_progress(client):
    r = client.post("/api/load", json={"tickers": ["^NSEI", "^CNXIT"], "period": "5Y"})
    assert r.status_code == 202
    job = client.app.state.ctx.jobs.wait(r.json()["id"])
    assert job["status"] == "done" and job["completed"] == 2 and job["progress"] == 1.0
    assert {i["ticker"]: i["status"] for i in job["items"]} == {"^NSEI": "done", "^CNXIT": "done"}
    assert job["items"][0]["rows"] > 1000 and len(job["items"][0]["first"]) == 10
    assert client.get(f"/api/jobs/{job['id']}").json()["status"] == "done"
    assert client.get("/api/jobs/missing").status_code == 404


def test_load_job_reports_per_ticker_errors(client):
    job_id = client.post("/api/load", json={"tickers": ["^NSEI", "NOPE"], "source": "csv"}).json()["id"]
    job = client.app.state.ctx.jobs.wait(job_id)
    assert job["status"] == "failed" and job["errors"] == 2
    assert "no CSV imported" in job["items"][1]["message"]


def test_load_validation(client):
    assert client.post("/api/load", json={"tickers": []}).status_code == 422
    assert client.post("/api/load", json={"tickers": [f"T{i}" for i in range(21)]}).status_code == 400


def test_job_sse_stream(client):
    job_id = client.post("/api/load", json={"tickers": ["^NSEI"], "period": "1Y"}).json()["id"]
    with client.stream("GET", f"/api/jobs/{job_id}/stream") as r:
        assert r.headers["content-type"].startswith("text/event-stream")
        text = "".join(r.iter_text())
    events = [block.split("\n") for block in text.strip().split("\n\n")]
    assert events[-1][0] == "event: done"
    assert json.loads(events[-1][1].removeprefix("data: "))["status"] == "done"
    assert all(e[0] in ("event: progress", "event: done") for e in events)


# --------------------------------------------------------------------------- #
# CSV import & cache
# --------------------------------------------------------------------------- #
CSV = "Date,Close\n2024-01-02,100\n2024-01-03,101.5\n2024-01-04,99.25\n"


def test_csv_import_then_series(client):
    r = client.post("/api/import/csv?ticker=MINE", content=CSV, headers={"Content-Type": "text/csv"})
    assert r.status_code == 201 and r.json() == {"ticker": "MINE", "source": "csv", "rows": 3,
                                                 "first": "2024-01-02", "last": "2024-01-04"}
    s = client.get("/api/series/MINE?source=csv&start=2024-01-01&columns=close,return").json()
    assert [row[1] for row in s["rows"]] == [100, 101.5, 99.25]
    assert client.post("/api/import/csv?ticker=MINE", content="").status_code == 400
    assert client.post("/api/import/csv?ticker=MINE", content="a,b\n1,2\n").status_code == 400
    assert client.post("/api/import/csv?ticker=bad%20name", content=CSV).status_code == 422


def test_cache_inventory_update_and_delete(client):
    client.get(f"/api/series/{NIFTY}?period=1Y")
    inv = client.get("/api/cache").json()
    assert_api_conventions(inv)
    assert [(e["source"], e["ticker"]) for e in inv["entries"]] == [("demo", "^NSEI")]
    assert inv["total_rows"] == inv["entries"][0]["rows"]

    job_id = client.post("/api/cache/update", json={}).json()["id"]
    assert client.app.state.ctx.jobs.wait(job_id)["status"] == "done"

    assert client.delete(f"/api/cache/demo/{NIFTY}").json() == {"deleted_files": 2}
    assert client.delete(f"/api/cache/demo/{NIFTY}").status_code == 404
    assert client.post("/api/cache/update", json={}).status_code == 400  # nothing cached


def test_delete_all_needs_confirm(client):
    client.get(f"/api/series/{NIFTY}?period=1Y")
    assert client.delete("/api/cache").status_code == 400
    assert client.delete("/api/cache?confirm=true").json()["deleted_files"] == 2


def test_ticker_check(client):
    r = client.get("/api/tickers/check").json()
    assert r["source"] == "demo" and r["checked"] >= 30 and r["ok"] == r["checked"]


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
EXPORT = {"tickers": ["^NSEI", "^NSEBANK"], "start": "2023-01-01", "end": "2023-12-31",
          "freq": "Monthly", "columns": ["close", "return", "sma_3"]}


def test_export_xlsx(client):
    r = client.post("/api/export", json={**EXPORT, "format": "xlsx"})
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('_monthly_' +
                                                                              pd.Timestamp.today().strftime("%Y%m%d") + '.xlsx"')
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Closes", "NIFTY 50", "NIFTY Bank", "About"]
    ws = wb["NIFTY 50"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Date", "Close", "Return", "SMA 3"]
    assert ws.cell(3, 3).number_format == "0.00%"          # returns are real % cells
    assert abs(ws.cell(3, 3).value) < 1                     # …holding decimals
    assert ws.max_row == 13                                 # header + 12 months
    assert ws["A2"].number_format == "dd-mm-yyyy"


def test_export_xlsx_options(client):
    client.patch("/api/settings/export", json={"metadata_sheet": False, "excel_freeze_panes": False})
    r = client.post("/api/export", json={**EXPORT, "format": "xlsx",
                                         "options": {"indian_number_format": True, "decimals": 1,
                                                     "date_format": "YYYY-MM-DD"}})
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert "About" not in wb.sheetnames
    ws = wb["NIFTY 50"]
    assert ws.freeze_panes is None
    assert ws.cell(2, 2).number_format.startswith("[>=10000000]") and ws.cell(2, 2).number_format.endswith(".0")
    assert ws["A2"].number_format == "yyyy-mm-dd"


def test_export_csv_long_format(client):
    r = client.post("/api/export", json={**EXPORT, "format": "csv"})
    assert r.headers["content-type"].startswith("text/csv")
    df = pd.read_csv(io.StringIO(r.content.decode("utf-8-sig")))
    assert list(df.columns) == ["Date", "Ticker", "Close", "Return", "SMA 3"]
    assert set(df["Ticker"]) == {"^NSEI", "^NSEBANK"} and len(df) == 24
    assert df["Date"].iloc[0] == "31-01-2023"               # settings date format
    assert (df["Close"] == df["Close"].round(2)).all()


def test_export_zip_and_json(client):
    z = zipfile.ZipFile(io.BytesIO(client.post("/api/export", json={**EXPORT, "format": "zip"}).content))
    assert sorted(z.namelist()) == ["NIFTY_50.csv", "NIFTY_Bank.csv"]
    j = json.loads(client.post("/api/export", json={**EXPORT, "format": "json"}).content)
    assert j["series"]["^NSEI"]["columns"] == ["date", "close", "return", "sma_3"]
    assert len(j["series"]["^NSEI"]["rows"]) == 12
    assert_api_conventions(j)


def test_export_preset(client):
    client.patch("/api/settings/export", json={"presets": [
        {"name": "Bank yearly", "tickers": ["^NSEBANK"], "frequency": "Yearly", "columns": ["close"], "format": "json"}]})
    j = json.loads(client.post("/api/export", json={"preset": "Bank yearly", "start": "2015-01-01",
                                                    "end": "2020-12-31"}).content)
    assert list(j["series"]) == ["^NSEBANK"] and j["freq"] == "Yearly"
    assert len(j["series"]["^NSEBANK"]["rows"]) == 6
    assert client.post("/api/export", json={"preset": "missing"}).status_code == 404
    assert client.post("/api/export", json={"format": "csv"}).status_code == 400   # no tickers


def test_no_nan_in_json(client):
    r = client.get(f"/api/series/{NIFTY}?period=1Y&columns=close,sma_200")
    assert "NaN" not in r.text
    assert np.isfinite([v for row in r.json()["rows"] for v in row[1:] if v is not None]).all()


def test_volume_available_follows_the_data():
    """Indices are assumed to have no volume, but if the source does send it the API says so."""
    import pandas as pd
    from api.routers.series import caveats, volume_available
    idx = pd.bdate_range("2024-01-01", periods=10)
    with_vol = pd.DataFrame({"Close": 1.0, "Volume": 1000.0}, index=idx)
    no_vol = pd.DataFrame({"Close": 1.0, "Volume": 0.0}, index=idx)
    assert volume_available("^NSEI", with_vol)
    assert not volume_available("^NSEI", no_vol)
    assert not volume_available("^NSEI")                 # no data: fall back to the catalogue
    assert volume_available("NIFTYBEES.NS")              # stocks/ETFs always have volume
    assert "Volume is not available for this series." not in caveats("^NSEI", "yahoo", with_vol)
    assert "Volume is not available for this series." in caveats("^NSEI", "yahoo", no_vol)
