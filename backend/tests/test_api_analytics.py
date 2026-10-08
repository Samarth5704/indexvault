"""API: analytics endpoints and SIP. Values are checked against core where cheap."""
import pandas as pd
import pytest

from conftest import assert_api_conventions
from indexvault import analytics as an
from indexvault import data

TWO = "tickers=^NSEI,^NSEBANK"
RANGE = "start=2010-01-01&end=2020-12-31"


def core_close(ticker="^NSEI", start="2010-01-01", end="2020-12-31"):
    return data.get_data(ticker, start, end, source="demo")["Close"]


def get(client, path):
    r = client.get(path)
    assert r.status_code == 200, r.text
    body = r.json()
    assert_api_conventions(body)
    return body


def test_summary_matches_core(client):
    body = get(client, f"/api/analytics/summary?{TWO}&{RANGE}")
    m = body["series"]["^NSEI"]["metrics"]
    core = an.summary_metrics(core_close())
    assert m["cagr"] == pytest.approx(core["CAGR %"] / 100)
    assert m["sharpe"] == pytest.approx(core["Sharpe"])
    assert m["max_dd_date"] == str(core["Max DD date"])
    assert body["series"]["^NSEI"]["start"] == "2010-01-01"
    assert body["kpi_cards"][0] == "end_level"
    assert {x["id"] for x in body["metrics"]} >= {"cagr", "ann_vol", "pct_from_52w_high"}


def test_summary_rf_override_and_settings(client):
    base = get(client, f"/api/analytics/summary?tickers=^NSEI&{RANGE}")["series"]["^NSEI"]["metrics"]["sharpe"]
    lower_rf = get(client, f"/api/analytics/summary?tickers=^NSEI&{RANGE}&rf=0")["series"]["^NSEI"]["metrics"]["sharpe"]
    assert lower_rf > base
    client.patch("/api/settings/analytics", json={"risk_free_rate": 0.0})
    assert get(client, f"/api/analytics/summary?tickers=^NSEI&{RANGE}")["params"]["rf"] == 0.0


def test_trailing(client):
    body = get(client, "/api/analytics/trailing?tickers=^NSEI&end=2020-12-31&periods=1M,1Y,3Y,YTD")
    assert body["periods"] == ["1M", "1Y", "3Y", "YTD"]
    core = an.trailing_returns(core_close("^NSEI", "2015-01-01"), ["1Y", "3Y"])
    assert body["series"]["^NSEI"]["1Y"] == pytest.approx(core["1Y"] / 100)
    assert body["series"]["^NSEI"]["3Y"] == pytest.approx(core["3Y"] / 100)
    assert client.get("/api/analytics/trailing?tickers=^NSEI&periods=2X").status_code == 400


def test_drawdowns(client):
    body = get(client, f"/api/analytics/drawdowns?ticker=^NSEI&{RANGE}&top=3")
    assert len(body["episodes"]) == 3
    e = body["episodes"][0]
    assert set(e) == {"peak", "trough", "recovered", "depth", "peak_to_trough_days",
                      "trough_to_recovery_days", "total_days", "ongoing"}
    assert -1 < e["depth"] < 0
    assert body["underwater"]["values"][0] == 0
    assert get(client, f"/api/analytics/drawdowns?ticker=^NSEI&{RANGE}")["episodes"].__len__() == 5  # settings default


def test_monthly_grid_and_yearly(client):
    grid = get(client, f"/api/analytics/monthly-grid?ticker=^NSEI&{RANGE}")
    assert grid["months"][0] == "Jan" and grid["years"] == list(range(2010, 2021))
    assert len(grid["values"][0]) == 12
    yearly = get(client, f"/api/analytics/yearly?ticker=^NSEI&{RANGE}")["years"]
    assert yearly[0]["year"] == 2010 and yearly[0]["start"] == "2010-01-01"
    assert yearly[3]["return"] == pytest.approx(grid["year_total"][3])
    c = core_close()
    assert yearly[3]["return"] == pytest.approx(c.loc["2013"].iloc[-1] / c.loc["2012"].iloc[-1] - 1)


def test_distribution(client):
    body = get(client, f"/api/analytics/distribution?ticker=^NSEI&{RANGE}&bins=20")
    assert len(body["edges"]) == 21 and len(body["counts"]) == 20 == len(body["normal"])
    assert sum(body["counts"]) == body["n"]
    assert sum(body["normal"]) == pytest.approx(body["n"], rel=0.05)
    assert len(get(client, f"/api/analytics/distribution?ticker=^NSEI&{RANGE}")["counts"]) == 50


def test_rolling_vol_is_warmed_up(client):
    body = get(client, "/api/analytics/rolling-vol?ticker=^NSEI&start=2020-01-01&end=2020-12-31&window=21")
    assert body["series"]["dates"][0] == "2020-01-01" and body["window"] == 21
    assert 0 < body["series"]["values"][0] < 2


def test_compare(client):
    body = get(client, f"/api/analytics/compare?{TWO}&{RANGE}&benchmark=^NSEI")
    assert body["rebased"]["series"]["^NSEI"][0] == pytest.approx(100)
    assert body["beta"]["^NSEI"] == pytest.approx(1.0)
    assert body["correlation"]["matrix"][0][0] == pytest.approx(1.0)
    assert list(body["metrics"]["^NSEI"]) == body["metric_ids"]
    assert [m["id"] for m in body["metric_meta"]] == body["metric_ids"] and body["metric_meta"][0]["kind"] == "pct"
    assert body["scatter"][1]["ticker"] == "^NSEBANK" and body["scatter"][1]["ann_vol"] > 0
    assert body["frequency"] == "W"
    monthly = get(client, f"/api/analytics/compare?{TWO}&{RANGE}&freq=M")
    assert monthly["correlation"]["matrix"][0][1] != body["correlation"]["matrix"][0][1]


def test_relative_strength_and_rolling_correlation(client):
    rs = get(client, f"/api/analytics/relative-strength?a=^NSEI&b=^NSEBANK&{RANGE}")
    assert rs["ratio"]["values"][0] == pytest.approx(100) and rs["sma_window"] == 50
    assert len(rs["sma"]["values"]) == len(rs["ratio"]["values"]) - 49
    rc = get(client, f"/api/analytics/rolling-correlation?a=^NSEI&b=^NSEBANK&{RANGE}&window=26")
    assert rc["window"] == 26 and all(-1 <= v <= 1 for v in rc["series"]["values"])


def test_rolling_returns_and_insight(client):
    body = get(client, "/api/analytics/rolling?tickers=^NSEI&years=1,5&target=0.1")
    five = body["series"]["^NSEI"]["windows"]["5"]
    rc = an.rolling_cagr(data.get_data("^NSEI", source="demo")["Close"], 5)
    assert five["observations"] == len(rc)
    assert five["median"] == pytest.approx(rc.median())
    assert five["pct_negative"] == pytest.approx((rc < 0).mean())
    assert five["pct_above_target"] == pytest.approx((rc > 0.1).mean())
    assert five["q1"] < five["median"] < five["q3"] and five["latest"] == pytest.approx(rc.iloc[-1])
    assert five["insight"].startswith("Over ") and "five-year periods, NIFTY 50" in five["insight"]
    assert "beat 10% a year" in five["insight"]
    assert client.get("/api/analytics/rolling?tickers=^NSEI&years=0").status_code == 400


def test_insight_number_grouping():
    from api.routers.compare import group_int
    assert group_int(1234567, "indian") == "12,34,567"
    assert group_int(1234567, "international") == "1,234,567"
    assert group_int(999, "indian") == "999" and group_int(-123456, "indian") == "-1,23,456"


def test_seasonality_and_weekday(client):
    s = get(client, f"/api/analytics/seasonality?ticker=^NSEI&{RANGE}&min_years=12")
    assert [m["month"] for m in s["months"]][:2] == ["Jan", "Feb"]
    assert set(s["months"][0]) == {"month", "avg", "median", "positive", "best", "worst", "years", "enough_data"}
    assert 0 <= s["months"][0]["positive"] <= 1 and s["months"][0]["enough_data"] is False
    w = get(client, f"/api/analytics/weekday?ticker=^NSEI&{RANGE}")
    assert [d["day"] for d in w["days"]] == ["Mon", "Tue", "Wed", "Thu", "Fri"]


def test_quality_and_big_moves(client):
    q = get(client, f"/api/analytics/quality?{TWO}&{RANGE}")
    nifty = q["series"]["^NSEI"]
    assert nifty["status"] in ("good", "warning", "critical")
    assert nifty["coverage"] == pytest.approx(1.0) and nifty["duplicate_dates"] == 0
    assert nifty["zero_volume_days"] is None                # index: volume not applicable
    moves = get(client, f"/api/analytics/big-moves?ticker=^NSEI&{RANGE}&threshold=0.04")["moves"]
    assert moves and all(abs(m["move"]) > 0.04 for m in moves)
    assert moves == sorted(moves, key=lambda m: m["date"])
    assert nifty["big_moves"] == len(get(client, f"/api/analytics/big-moves?ticker=^NSEI&{RANGE}")["moves"])


def test_sip(client):
    r = client.post("/api/analytics/sip", json={"ticker": "^NSEI", "start": "2015-01-01", "end": "2020-12-31",
                                                "amount": 5000, "step_up": 0.1,
                                                "top_ups": [{"date": "2018-03-10", "amount": 20000}]})
    assert r.status_code == 200, r.text
    body = r.json()
    assert_api_conventions(body)
    st = body["stats"]
    _, core = an.sip_backtest(core_close("^NSEI", "2015-01-01"), amount=5000, day_of_month=5,
                              step_up_pct=10, top_ups=[("2018-03-10", 20000)])
    assert st["xirr"] == pytest.approx(core["SIP XIRR %"] / 100)
    assert st["total_invested"] == pytest.approx(core["Total invested"]) and st["top_ups"] == 1
    assert body["params"]["day"] == 5                       # settings default
    assert body["ledger"]["columns"] == ["date", "price", "invested", "units", "value", "gain", "lump_sum_value"]
    assert len(body["ledger"]["rows"]) == 72
    assert "sensitivity" not in body


def test_sip_export(client):
    import io
    from openpyxl import load_workbook
    body = {"ticker": "^NSEI", "start": "2015-01-01", "end": "2020-12-31", "amount": 5000, "step_up": 0.1}
    r = client.post("/api/analytics/sip/export?format=xlsx", json=body)
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Ledger"]
    summary = {c.value: wb["Summary"].cell(2, i + 1) for i, c in enumerate(wb["Summary"][1])}
    assert summary["XIRR"].number_format == "0.00%" and abs(summary["XIRR"].value) < 1
    assert summary["Annual step-up"].value == pytest.approx(0.1)
    ledger = wb["Ledger"]
    assert [c.value for c in ledger[1]] == ["Date", "Price", "Invested", "Units", "Value", "Gain", "Lump-sum value"]
    assert ledger.max_row == 73                                  # header + 72 months
    csv = client.post("/api/analytics/sip/export?format=csv", json=body)
    assert csv.headers["content-type"].startswith("text/csv") and csv.content.decode("utf-8-sig").startswith("Date,Price")
    assert client.post("/api/analytics/sip/export?format=pdf", json=body).status_code == 422


def test_sip_sensitivity_and_errors(client):
    body = client.post("/api/analytics/sip?sensitivity=true",
                       json={"ticker": "^NSEI", "start": "2015-01-01", "end": "2020-12-31"}).json()
    sens = body["sensitivity"]
    assert sens[0]["start"].startswith("2015-01") and sens[0]["months"] == 72
    assert len(sens) == 72 - 12 + 1
    assert sens[0]["xirr"] == pytest.approx(body["stats"]["xirr"])
    r = client.post("/api/analytics/sip", json={"ticker": "^NSEI", "start": "2020-01-01", "end": "2020-01-10"})
    assert r.status_code == 400
    assert client.post("/api/analytics/sip", json={"ticker": "^NSEI", "step_up": 10}).status_code == 422


def test_memoised_summary_recomputes_after_cache_update(client):
    path = f"/api/analytics/summary?tickers=^NSEI&{RANGE}"
    first = get(client, path)["series"]["^NSEI"]["metrics"]["cagr"]
    memo = client.app.state.ctx.memo
    hits = memo.hits
    assert get(client, path)["series"]["^NSEI"]["metrics"]["cagr"] == first
    assert memo.hits > hits
    data.clear_cache("^NSEI", "demo")                       # cache file gone -> new mtime key
    assert get(client, path)["series"]["^NSEI"]["metrics"]["cagr"] == pytest.approx(first)


@pytest.mark.parametrize("path", [
    "/api/analytics/summary", "/api/analytics/drawdowns", "/api/analytics/compare?tickers=",
    "/api/analytics/relative-strength?a=^NSEI",
])
def test_missing_params(client, path):
    r = client.get(path)
    assert r.status_code in (400, 422) and "error" in r.json()


def test_too_many_tickers(client):
    many = ",".join(f"T{i}" for i in range(21))
    r = client.get(f"/api/analytics/summary?tickers={many}")
    assert r.status_code == 400 and r.json()["error"]["detail"]["given"] == 21


def test_unknown_ticker_404(client):
    # demo makes up data for any ticker, so use the csv source for a true miss
    r = client.get("/api/analytics/drawdowns?ticker=NOPE&source=csv")
    assert r.status_code == 404


def test_period_defaults_from_settings(client):
    client.patch("/api/settings/data", json={"default_period": "3Y"})
    body = get(client, "/api/analytics/drawdowns?ticker=^NSEI")
    start = pd.Timestamp(body["start"])
    assert pd.Timestamp.today() - pd.DateOffset(years=3, days=5) <= start <= pd.Timestamp.today() - pd.DateOffset(years=3) + pd.DateOffset(days=5)
