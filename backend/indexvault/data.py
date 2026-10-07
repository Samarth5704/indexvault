"""Data layer: fetch from Yahoo Finance, cache locally, update incrementally,
resample to other frequencies and export.

Cache layout (one CSV + one JSON per ticker):
    data_cache/yahoo/<SAFE_TICKER>.csv
    data_cache/yahoo/<SAFE_TICKER>.json   -> {"ticker", "requested_start", "last_update"}
"""
from __future__ import annotations

import io
import json
import os
import re
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

COLUMNS = ["Open", "High", "Low", "Close", "Adj Close", "Volume"]
EARLIEST = pd.Timestamp("1990-01-01")

CACHE_ROOT = Path(
    os.environ.get(
        "INDEXVAULT_CACHE", Path(__file__).resolve().parent.parent / "data_cache"
    )
)

FREQUENCIES = {
    "Daily": None,
    "Weekly": "W-FRI",
    "Monthly": "ME",
    "Quarterly": "QE",
    "Yearly": "YE",
}


# --------------------------------------------------------------------------- #
# Cache helpers
# --------------------------------------------------------------------------- #
def _safe(ticker: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]", "_", ticker)


def _paths(ticker: str, source: str) -> tuple[Path, Path]:
    d = CACHE_ROOT / source
    d.mkdir(parents=True, exist_ok=True)
    s = _safe(ticker)
    return d / f"{s}.csv", d / f"{s}.json"


def load_cached(ticker: str, source: str = "yahoo") -> tuple[pd.DataFrame | None, dict]:
    csv, meta = _paths(ticker, source)
    if not csv.exists():
        return None, {}
    df = pd.read_csv(csv, index_col=0, parse_dates=True)
    df.index.name = "Date"
    info = json.loads(meta.read_text()) if meta.exists() else {}
    return df, info


def _save(ticker: str, source: str, df: pd.DataFrame, info: dict) -> None:
    csv, meta = _paths(ticker, source)
    df.to_csv(csv, float_format="%.4f")
    info = {**info, "ticker": ticker, "last_update": datetime.now().isoformat(timespec="seconds")}
    meta.write_text(json.dumps(info, indent=2))


def cache_inventory() -> pd.DataFrame:
    """One row per cached ticker across all sources."""
    rows = []
    if CACHE_ROOT.exists():
        for src_dir in sorted(p for p in CACHE_ROOT.iterdir() if p.is_dir()):
            for meta in sorted(src_dir.glob("*.json")):
                info = json.loads(meta.read_text())
                csv = meta.with_suffix(".csv")
                if not csv.exists():
                    continue
                idx = pd.read_csv(csv, usecols=[0], index_col=0, parse_dates=True).index
                rows.append({
                    "Source": src_dir.name,
                    "Ticker": info.get("ticker", meta.stem),
                    "Rows": len(idx),
                    "First date": idx.min().date() if len(idx) else None,
                    "Last date": idx.max().date() if len(idx) else None,
                    "Last updated": info.get("last_update", ""),
                    "Size (KB)": round(csv.stat().st_size / 1024, 1),
                })
    return pd.DataFrame(rows)


def clear_cache(ticker: str | None = None, source: str | None = None) -> int:
    n = 0
    if not CACHE_ROOT.exists():
        return 0
    for src_dir in CACHE_ROOT.iterdir():
        if not src_dir.is_dir() or (source and src_dir.name != source):
            continue
        pattern = f"{_safe(ticker)}.*" if ticker else "*"
        for f in src_dir.glob(pattern):
            if f.suffix in (".csv", ".json"):
                f.unlink()
                n += 1
    return n


# --------------------------------------------------------------------------- #
# Sources
# --------------------------------------------------------------------------- #
def _clean(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(columns=COLUMNS)
    if isinstance(df.columns, pd.MultiIndex):
        df = df.droplevel(-1, axis=1) if df.columns.nlevels > 1 else df
    df = df.copy()
    if "Adj Close" not in df.columns and "Close" in df.columns:
        df["Adj Close"] = df["Close"]
    if "Volume" not in df.columns:
        df["Volume"] = 0
    df = df[[c for c in COLUMNS if c in df.columns]]
    idx = pd.to_datetime(df.index)
    if getattr(idx, "tz", None) is not None:
        idx = idx.tz_localize(None)
    df.index = idx.normalize()
    df.index.name = "Date"
    df = df[~df.index.duplicated(keep="last")].sort_index()
    df = df.dropna(subset=["Close"])
    df = df[df["Close"] > 0]
    return df


def _fetch_yahoo(ticker: str, start: pd.Timestamp, end: pd.Timestamp) -> pd.DataFrame:
    import yfinance as yf

    raw = yf.Ticker(ticker).history(
        start=start.strftime("%Y-%m-%d"),
        end=(end + timedelta(days=1)).strftime("%Y-%m-%d"),
        interval="1d",
        auto_adjust=False,
        actions=False,
    )
    return _clean(raw)


def _fetch_demo(ticker: str, start: pd.Timestamp, end: pd.Timestamp) -> pd.DataFrame:
    """Deterministic synthetic OHLCV (regime-switching random walk) for offline use."""
    seed = sum(ord(c) * (i + 1) for i, c in enumerate(ticker)) % (2**32)
    rng = np.random.default_rng(seed)
    days = pd.bdate_range("2000-01-03", pd.Timestamp.today().normalize())
    n = len(days)
    drift = rng.uniform(0.07, 0.16) / 252
    base_vol = rng.uniform(0.14, 0.26) / np.sqrt(252)
    # volatility regimes: calm / stressed
    regime = np.zeros(n)
    state = 0
    for i in range(n):
        if rng.random() < (0.004 if state == 0 else 0.02):
            state = 1 - state
        regime[i] = state
    vol = base_vol * np.where(regime == 1, 2.2, 1.0)
    rets = drift - np.where(regime == 1, 0.0004, 0) + vol * rng.standard_t(5, n) / np.sqrt(5 / 3)
    close = rng.uniform(800, 6000) * np.exp(np.cumsum(rets))
    open_ = close * np.exp(rng.normal(0, vol * 0.3))
    high = np.maximum(open_, close) * np.exp(np.abs(rng.normal(0, vol * 0.5)))
    low = np.minimum(open_, close) * np.exp(-np.abs(rng.normal(0, vol * 0.5)))
    volume = 0 if ticker.startswith("^") else rng.integers(1e5, 5e6, n)
    df = pd.DataFrame(
        {"Open": open_, "High": high, "Low": low, "Close": close, "Adj Close": close, "Volume": volume},
        index=days,
    )
    return _clean(df.loc[start:end])


SOURCES = {"yahoo": _fetch_yahoo, "demo": _fetch_demo}


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #
def get_data(
    ticker: str,
    start: str | date | None = None,
    end: str | date | None = None,
    source: str = "yahoo",
    refresh: bool = False,
    use_cache: bool = True,
) -> pd.DataFrame:
    """Daily OHLCV for `ticker` between start and end (inclusive).

    Uses the local cache when possible and only downloads what is missing:
      * newer rows since the last cached date (incremental update)
      * older rows if you ask for an earlier start than ever before (backfill)
    `refresh=True` re-downloads the full requested range.
    """
    fetch = SOURCES[source]
    start = pd.Timestamp(start) if start else EARLIEST
    end = pd.Timestamp(end) if end else pd.Timestamp.today().normalize()

    if not use_cache:
        return fetch(ticker, start, end)

    cached, info = load_cached(ticker, source)
    if refresh or cached is None or cached.empty:
        df = fetch(ticker, start, end)
        if not df.empty:
            _save(ticker, source, df, {"requested_start": str(start.date())})
        return df.loc[start:end]

    pieces = [cached]
    req_start = pd.Timestamp(info.get("requested_start", cached.index.min()))
    if start < req_start:  # backfill older history
        older = fetch(ticker, start, cached.index.min() - timedelta(days=1))
        pieces.insert(0, older)
        req_start = start
    last = cached.index.max()
    today = pd.Timestamp.today().normalize()
    try:
        since_update = datetime.now() - datetime.fromisoformat(info.get("last_update", "1990-01-01"))
    except ValueError:
        since_update = timedelta(days=999)
    stale = last < today - pd.offsets.BDay(3) or since_update > timedelta(hours=3)
    if end > last and last < today and stale:
        # small overlap so revised last bars get corrected
        newer = fetch(ticker, last - timedelta(days=5), today)
        pieces.append(newer)

    if len(pieces) > 1:
        merged = pd.concat([p for p in pieces if not p.empty])
        merged = merged[~merged.index.duplicated(keep="last")].sort_index()
        _save(ticker, source, merged, {"requested_start": str(req_start.date())})
        cached = merged
    return cached.loc[start:end]


def check_ticker(ticker: str, source: str = "yahoo") -> tuple[bool, str]:
    """Quick probe: does the source return recent data for this ticker?"""
    try:
        end = pd.Timestamp.today().normalize()
        df = SOURCES[source](ticker, end - timedelta(days=20), end)
        if df.empty:
            return False, "no data returned"
        return True, f"last close {df['Close'].iloc[-1]:,.2f} on {df.index[-1].date()}"
    except Exception as e:  # noqa: BLE001
        return False, str(e)[:120]


def resample(df: pd.DataFrame, freq: str, add_returns: bool = True) -> pd.DataFrame:
    """Aggregate daily OHLCV to Weekly/Monthly/Quarterly/Yearly bars.

    Each bar is labelled with the *last actual trading date* inside it
    (not a calendar month-end that may be a holiday).
    """
    rule = FREQUENCIES[freq]
    out = df.copy()
    if rule is not None and not df.empty:
        g = df.resample(rule)
        out = g.agg({
            "Open": "first", "High": "max", "Low": "min",
            "Close": "last", "Adj Close": "last", "Volume": "sum",
        })
        last_dates = df.index.to_series().resample(rule).max()
        out.index = last_dates.reindex(out.index).values
        out = out.dropna(subset=["Close"])
        out.index = pd.DatetimeIndex(out.index, name="Date")
    if add_returns and not out.empty:
        out["Return %"] = out["Close"].pct_change() * 100
    return out


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
def _sheet_name(name: str, used: set[str]) -> str:
    s = re.sub(r"[\[\]\*\?/\\:]", "", name)[:31] or "Sheet"
    base, i = s, 2
    while s in used:
        s = f"{base[:28]}_{i}"
        i += 1
    used.add(s)
    return s


def to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Write several DataFrames to one formatted .xlsx workbook."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            sname = _sheet_name(name, used)
            frame = frame.copy()
            if isinstance(frame.index, pd.DatetimeIndex):
                frame.index = frame.index.date
            frame.to_excel(xw, sheet_name=sname)
            ws = xw.sheets[sname]
            ws.freeze_panes = "B2"
            head_fill = PatternFill("solid", fgColor="1C5CAB")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center")
            for col_idx, col in enumerate(ws.columns, start=1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 40)
                for c in list(col)[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "#,##0.00"
    return buf.getvalue()


def to_csv_zip_bytes(frames: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, frame in frames.items():
            z.writestr(f"{_safe(name)}.csv", frame.to_csv(float_format="%.4f"))
    return buf.getvalue()


def combine_close(frames: dict[str, pd.DataFrame], column: str = "Close") -> pd.DataFrame:
    """Wide table: one column per ticker, aligned on dates (outer join)."""
    return pd.DataFrame({k: v[column] for k, v in frames.items() if not v.empty}).sort_index()
