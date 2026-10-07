"""Excel and CSV-zip writers for exported tables."""
from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime

import pandas as pd


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]", "_", name)


def _sheet_name(name: str, used: set[str]) -> str:
    s = re.sub(r"[\[\]\*\?/\\:]", "", name)[:31] or "Sheet"
    base, i = s, 2
    while s in used:
        s = f"{base[:28]}_{i}"
        i += 1
    used.add(s)
    return s


def to_excel_bytes(
    sheets: dict[str, pd.DataFrame],
    header_colour: str = "1C5CAB",
    freeze_panes: bool = True,
    number_format: str = "#,##0.00",
    date_format: str | None = None,
    column_formats: dict[str, str] | None = None,
) -> bytes:
    """Write several DataFrames to one formatted .xlsx workbook.

    `header_colour` is a hex fill (with or without '#'); `number_format` and
    `date_format` are Excel format codes (e.g. "dd-mm-yyyy") applied to float
    and date cells. `column_formats` maps a column header to its own format
    code (e.g. {"Return": "0.00%"}), overriding `number_format`."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            sname = _sheet_name(name, used)
            frame = frame.copy()
            if isinstance(frame.index, pd.DatetimeIndex):
                frame.index = pd.Index(frame.index.date, name=frame.index.name)
            frame.to_excel(xw, sheet_name=sname)
            ws = xw.sheets[sname]
            if freeze_panes:
                ws.freeze_panes = "B2"
            head_fill = PatternFill("solid", fgColor=header_colour.lstrip("#").upper())
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center")
            for col_idx, col in enumerate(ws.columns, start=1):
                col_format = (column_formats or {}).get(str(col[0].value), number_format)
                width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 40)
                for c in list(col)[1:]:
                    if isinstance(c.value, float):
                        c.number_format = col_format
                    elif date_format and isinstance(c.value, (date, datetime)):
                        c.number_format = date_format
    return buf.getvalue()


def to_csv_zip_bytes(frames: dict[str, pd.DataFrame], float_format: str = "%.4f",
                     date_format: str | None = None) -> bytes:
    """One CSV per frame in a zip. `date_format` is a strftime pattern."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, frame in frames.items():
            z.writestr(f"{_safe_name(name)}.csv",
                       frame.to_csv(float_format=float_format, date_format=date_format))
    return buf.getvalue()
